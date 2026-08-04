#!/usr/bin/env python3
"""Extracts data.json for the Sub-10 dashboard from the Sistema de gestao spreadsheet.

Usage: python3 build_data.py <caminho-para-o-xlsx> [saida.json]

Regras aplicadas (decisoes tomadas na conversa com o treinador):
- Presenca de um atleta num treino = ele tem nota preenchida em "Avaliacao
  individual do treino" naquele dia (blank = ausencia).
- Referencia de minutagem por sessao = 80 minutos (confirmado, nao 90).
- Conteudo principal / secundario / concorrente sao papeis pedagogicos
  distintos (todos intencionais) e sao mantidos separados, nunca somados
  num unico numero.
- Tendencia do atleta (ultimos 5 treinos vs media geral): limiar de +/-0.5.
"""
import json
import re
import sys
import unicodedata
from datetime import datetime, date
from pathlib import Path

import openpyxl

PHOTOS_DIR = Path(__file__).parent / 'photos'


def slug(name):
    s = unicodedata.normalize('NFKD', str(name)).encode('ascii', 'ignore').decode()
    s = re.sub(r'[^a-zA-Z0-9]+', '-', s).strip('-').lower()
    return s


def norm_header(s):
    return ' '.join(str(s).split())


def header_map(ws, header_row_idx):
    """Maps normalized header text -> column index, reading a specific row.

    Looks up columns by name instead of position so that inserting, removing,
    or reordering a column in the spreadsheet doesn't silently break parsing
    (this bit us once already: a stray blank column got deleted upstream and
    every field after it shifted by one).
    """
    row = [c.value for c in ws[header_row_idx]]
    m = {}
    for idx, name in enumerate(row):
        if name:
            key = norm_header(name)
            if key not in m:
                m[key] = idx
    return m


def g(row, colmap, name, default=None):
    idx = colmap.get(name)
    if idx is None or idx >= len(row):
        return default
    return row[idx]

MONTH_PT = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez']
QUARTILE_ORDER = ['Q1', 'Q2', 'Q3', 'Q4']
TENURE_ORDER = ['0–3 meses', '4–6 meses', '7–12 meses', '13–24 meses', '+24 meses']
TASK_ORDER = ['Analítico geral', 'Analítico em contexto específico', 'Jogo fundamental', 'Conceitual', 'Situacional', 'Jogo condicionado', 'Jogo formal']
COMPLEX_ORDER = ['Baixa', 'Média', 'Alta']
PART_ORDER = ['Preparatória', 'Conceitual', 'Conexão']
TREND_THRESHOLD = 0.5
REFERENCE_MINUTES = 80


def d(dt):
    if dt is None:
        return None
    if isinstance(dt, datetime):
        dt = dt.date()
    return dt.isoformat()


def month_label(dt):
    if dt is None:
        return None
    if isinstance(dt, datetime):
        dt = dt.date()
    return f"{MONTH_PT[dt.month - 1]}/{str(dt.year)[2:]}"


def tenure_band(months):
    if months is None:
        return None
    if months <= 3:
        return '0–3 meses'
    if months <= 6:
        return '4–6 meses'
    if months <= 12:
        return '7–12 meses'
    if months <= 24:
        return '13–24 meses'
    return '+24 meses'


def months_between(start, end):
    if start is None:
        return None
    if isinstance(start, datetime):
        start = start.date()
    if isinstance(end, datetime):
        end = end.date()
    return (end.year - start.year) * 12 + (end.month - start.month) - (1 if end.day < start.day else 0)


def norm(s):
    return str(s).strip() if s is not None else ''


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else 'data.json'

    wb = openpyxl.load_workbook(src, data_only=True)
    as_of = date.today().isoformat()

    # ---------- Banco de atletas ----------
    ws = wb['Banco de atletas']
    col_a = header_map(ws, 1)
    athletes = []
    athlete_by_nick = {}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        name = g(row, col_a, 'Nome')
        if not name:
            continue
        nickname = g(row, col_a, 'Apelido')
        aid = g(row, col_a, 'ID atleta')
        dob = g(row, col_a, 'Data de Nascimento')
        quartile = g(row, col_a, 'Quartil')
        position = g(row, col_a, 'Posição')
        foot = g(row, col_a, 'Pé dominante')
        joined = g(row, col_a, 'Chegada ao clube')
        photo_link = g(row, col_a, 'Link foto')
        characteristic = g(row, col_a, 'Característica')
        improvement = g(row, col_a, 'Pontos a melhorar')
        tenure_months = months_between(joined, date.today()) if joined else None
        local_photo = next((PHOTOS_DIR / f"{slug(nickname)}{ext}" for ext in ('.jpg', '.jpeg', '.png')
                             if (PHOTOS_DIR / f"{slug(nickname)}{ext}").exists()), None)
        photo_external = (photo_link if photo_link and str(photo_link).startswith('http')
                           and 'drive.google.com' not in str(photo_link) else None)
        if local_photo:
            photo = f"./photos/{local_photo.name}"
        elif photo_external:
            photo = photo_external
        else:
            photo = None
        a = {
            'id': aid,
            'name': name,
            'nickname': nickname,
            'dob': d(dob),
            'birthMonth': month_label(dob) if dob else None,
            'quartile': quartile,
            'position': position,
            'foot': foot,
            'joined': d(joined),
            'tenureMonths': tenure_months,
            'tenureBand': tenure_band(tenure_months),
            'category': g(row, col_a, 'Categoria'),
            'photo': photo,
            'photoExternal': photo_external,
            'characteristic': characteristic,
            'improvement': improvement,
            'status': 'ativo',
        }
        athletes.append(a)
        athlete_by_nick[norm(nickname).lower()] = a

    # ---------- Banco de treinos ----------
    ws = wb['Banco de treinos']
    col_t = header_map(ws, 3)
    trainings = []
    blocks = []
    contents = []  # long format, role-tagged: principal / secundario / concorrente
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        dt = g(row, col_t, 'Data')
        session_id = g(row, col_t, 'ID da sessão')
        category = g(row, col_t, 'Categoria')
        c_prin_prep = g(row, col_t, 'Conteúdo principal - preparatória')
        c_sec_prep = g(row, col_t, 'Conteúdo secundário - preparatória')
        c_conc_prep = g(row, col_t, 'Conteúdo concorrente - preparatória')
        tipo_prep = g(row, col_t, 'Tipo de conteúdo - preparatória')
        tarefa_prep = g(row, col_t, 'Tipo de tarefa - preparatória')
        complex_prep = g(row, col_t, 'Complexidade - preparatória')
        c_prin_conc = g(row, col_t, 'Conteúdo principal - conceitual')
        c_sec_conc = g(row, col_t, 'Conteúdo secundário - conceitual')
        c_conc_conc = g(row, col_t, 'Conteúdo concorrente - conceitual')
        tipo_conc = g(row, col_t, 'Tipo de conteúdo - conceitual')
        tarefa_conc = g(row, col_t, 'Tipo de tarefa - conceitual')
        complex_conc = g(row, col_t, 'Complexidade - conceitual')
        relacao_conexao = g(row, col_t, 'Relação númerica - conexão')
        tarefa_conexao = g(row, col_t, 'Tipo de tarefa - conexão')
        complex_conexao = g(row, col_t, 'Complexidade - conexão')
        nota_geral = g(row, col_t, 'Nota geral da sessão (0 a 10)')
        num_atletas = g(row, col_t, 'Nº de atletas')
        atletas_avaliacao = g(row, col_t, 'Atletas em avaliação')
        min_prep = g(row, col_t, 'Minutagem - preparatória')
        jog_ativos_prep = g(row, col_t, 'Jogadores ativos - preparatória')
        taa_prep = g(row, col_t, 'TAA - preparatória')
        desemp_prep = g(row, col_t, 'Desempenho - preparatória (0 a 10)')
        min_conc = g(row, col_t, 'Minutagem - conceitual')
        jog_ativos_conc = g(row, col_t, 'Jogadores ativos - conceitual')
        taa_conc = g(row, col_t, 'TAA - conceitual')
        desemp_conc = g(row, col_t, 'Desempenho - conceitual (0 a 10)')
        min_conexao = g(row, col_t, 'Minutagem - conexão')
        jog_ativos_conexao = g(row, col_t, 'Jogadores ativos - conexão')
        taa_conexao = g(row, col_t, 'TAA - conexão')
        desemp_conexao = g(row, col_t, 'Desempenho - conexão (0 a 10)')
        observacoes = g(row, col_t, 'Observações')
        minutagem_efetiva = g(row, col_t, 'Minutagem efetiva')
        total_taa = g(row, col_t, 'Total TAA')
        pct_minutagem = g(row, col_t, '% de minutagem')
        pct_taa = g(row, col_t, '% de TAA')
        if dt is None:
            continue  # sessoes sem data (pos-treino parcial, ainda nao detalhadas)
        month = month_label(dt)
        trainings.append({
            'date': d(dt), 'month': month, 'sessionId': session_id, 'category': category,
            'note': nota_geral, 'numAthletes': num_atletas, 'athletesEval': atletas_avaliacao,
            'effectiveMinutes': minutagem_efetiva, 'totalTAA': total_taa,
            'pctMinutagem': pct_minutagem if isinstance(pct_minutagem, (int, float)) else None,
            'pctTAA': pct_taa if isinstance(pct_taa, (int, float)) else None,
            'observations': observacoes,
        })

        parts = [
            ('Preparatória', c_prin_prep, c_sec_prep, c_conc_prep, tipo_prep, tarefa_prep, complex_prep,
             min_prep, jog_ativos_prep, taa_prep, desemp_prep, None),
            ('Conceitual', c_prin_conc, c_sec_conc, c_conc_conc, tipo_conc, tarefa_conc, complex_conc,
             min_conc, jog_ativos_conc, taa_conc, desemp_conc, None),
            ('Conexão', None, None, None, 'Conexão com jogo', tarefa_conexao, complex_conexao,
             min_conexao, jog_ativos_conexao, taa_conexao, desemp_conexao, relacao_conexao),
        ]
        for part, c_prin, c_sec, c_conc, tipo, tarefa, complexidade, minutes, jog_ativos, taa, desemp, relacao in parts:
            if minutes is None and c_prin is None and relacao is None:
                continue
            blocks.append({
                'date': d(dt), 'month': month, 'sessionId': session_id, 'part': part,
                'content': relacao if part == 'Conexão' else c_prin,
                'contentType': tipo, 'taskType': tarefa, 'complexity': complexidade,
                'minutes': minutes, 'numAthletesActive': jog_ativos, 'taa': taa,
                'performance': desemp, 'sessionNote': nota_geral,
            })
            for role, content in (('principal', c_prin), ('secundario', c_sec), ('concorrente', c_conc)):
                if content:
                    contents.append({
                        'date': d(dt), 'month': month, 'part': part, 'role': role,
                        'content': content, 'contentType': tipo,
                    })

    # ---------- Avaliacao individual do treino ----------
    ws = wb['Avaliação individual do treino']
    header = [c.value for c in ws[1]]
    date_cols = header[2:]
    notes = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        nickname = row[0]
        if not nickname:
            continue
        a = athlete_by_nick.get(norm(nickname).lower())
        for col_idx, raw_date in enumerate(date_cols, start=2):
            note = row[col_idx]
            if note is None:
                continue
            dt = datetime.strptime(raw_date, '%d/%m/%Y').date()
            notes.append({
                'athlete': nickname, 'date': dt.isoformat(), 'month': month_label(dt), 'note': note,
                'position': a['position'] if a else None,
                'quartile': a['quartile'] if a else None,
                'foot': a['foot'] if a else None,
                'tenureBand': a['tenureBand'] if a else None,
                'photo': a['photo'] if a else None,
            })

    # ---------- Pos-jogo ----------
    games = []
    if 'Pós-jogo' in wb.sheetnames:
        ws = wb['Pós-jogo']
        col_g = header_map(ws, 2)
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
            game_no = g(row, col_g, 'Nº Jogo')
            if game_no is None:
                continue
            dt = g(row, col_g, 'Data')
            gf = g(row, col_g, 'Gols feitos')
            ga = g(row, col_g, 'Gols sofridos')
            result = g(row, col_g, 'Resultado')
            games.append({
                'gameNo': game_no, 'date': d(dt), 'month': month_label(dt), 'category': g(row, col_g, 'Categoria'),
                'opponent': g(row, col_g, 'Adversário'), 'rank': g(row, col_g, 'Ranking'),
                'oppCategory': g(row, col_g, 'Categoria adversário'), 'competition': g(row, col_g, 'Competição'),
                'local': g(row, col_g, 'Local'), 'city': g(row, col_g, 'Cidade'), 'period': g(row, col_g, 'Período'),
                'format': g(row, col_g, 'Formato'), 'platform': g(row, col_g, 'Plataforma de jogo'),
                'coach': g(row, col_g, 'Treinador'), 'captain': g(row, col_g, 'Capitão'),
                'minutes': g(row, col_g, 'Minutagem'), 'goalsFor': gf, 'goalsAgainst': ga,
                'saldo': g(row, col_g, 'Saldo'),
                'result': result if (gf is not None and ga is not None) else None,
                'yellow': g(row, col_g, 'Cartão amarelo'), 'red': g(row, col_g, 'Cartão vermelho'),
                'performance': g(row, col_g, 'Desempenho'), 'comments': g(row, col_g, 'Comentários gerais'),
                'positives': g(row, col_g, 'Pontos positivos'), 'negatives': g(row, col_g, 'Pontos negativos'),
                'daysSinceLastGame': g(row, col_g, 'Data entre jogos'),
            })

    # ---------- Pos-jogo individual ----------
    game_ind = []
    if 'Pós-jogo individual' in wb.sheetnames:
        ws = wb['Pós-jogo individual']
        col_gi = header_map(ws, 2)
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
            name = g(row, col_gi, 'Nome')
            if name is None:
                continue
            dt = g(row, col_gi, 'Data')
            position = g(row, col_gi, 'Posição')
            minutes = g(row, col_gi, 'Minutagem')
            goals = g(row, col_gi, 'Gols feitos')
            goals_against = g(row, col_gi, 'Gols sofridos')
            assists = g(row, col_gi, 'Assistências')
            yellow = g(row, col_gi, 'Cartão amarelo')
            red = g(row, col_gi, 'Cartão vermelho')
            a = athlete_by_nick.get(norm(name).lower())
            game_ind.append({
                'gameNo': g(row, col_gi, 'Nº Jogo'), 'date': d(dt), 'month': month_label(dt),
                'category': g(row, col_gi, 'Categoria'),
                'athlete': name, 'position': position or (a['position'] if a else None),
                'quartile': a['quartile'] if a else None, 'foot': a['foot'] if a else None,
                'titularity': g(row, col_gi, 'Titularidade'), 'minutes': minutes or 0,
                'performance': g(row, col_gi, 'Desempenho'),
                'goals': goals or 0, 'goalsAgainst': goals_against or 0, 'assists': assists or 0,
                'yellow': yellow or 0, 'red': red or 0, 'comments': g(row, col_gi, 'Comentários individuais'),
            })

    all_dates = [t['date'] for t in trainings] + [g['date'] for g in games]
    all_dated = trainings + games
    months_seen = sorted({t['month'] for t in all_dated if t['month']},
                          key=lambda m: min(t['date'] for t in all_dated if t['month'] == m))

    data = {
        'meta': {
            'generatedAt': datetime.now().isoformat(timespec='seconds'),
            'asOf': as_of,
            'minDate': min(all_dates) if all_dates else None,
            'maxDate': max(all_dates) if all_dates else None,
            'referenceMinutes': REFERENCE_MINUTES,
            'trendThreshold': TREND_THRESHOLD,
            'months': months_seen,
            'athletes': [a['nickname'] for a in athletes],
            'positions': sorted({a['position'] for a in athletes if a['position']}),
            'quartiles': QUARTILE_ORDER,
            'competitions': sorted({g['competition'] for g in games if g['competition']}),
            'ranks': sorted({g['rank'] for g in games if g['rank']}),
            'formats': sorted({g['format'] for g in games if g['format']}),
            'results': sorted({g['result'] for g in games if g['result']}),
            'taskTypes': sorted({bl['taskType'] for bl in blocks if bl['taskType']}),
            'complexities': sorted({bl['complexity'] for bl in blocks if bl['complexity']}),
            'parts': PART_ORDER,
            'contents': sorted({c['content'] for c in contents if c['content']}),
            'quartileOrder': QUARTILE_ORDER,
            'tenureOrder': TENURE_ORDER,
            'taskOrder': TASK_ORDER,
            'complexityOrder': COMPLEX_ORDER,
            'partOrder': PART_ORDER,
        },
        'athletes': athletes,
        'trainings': trainings,
        'blocks': blocks,
        'contents': contents,
        'notes': notes,
        'games': games,
        'gameInd': game_ind,
    }

    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=None, separators=(',', ':'))

    print(f"OK: {out_path}")
    print(f"  athletes={len(athletes)} trainings={len(trainings)} blocks={len(blocks)} "
          f"contents={len(contents)} notes={len(notes)} games={len(games)} gameInd={len(game_ind)}")


if __name__ == '__main__':
    main()
