#!/usr/bin/env python3
"""Baixa e re-hospeda as fotos dos atletas dentro do repositorio.

Por que existe: as fotos hoje estao linkadas em i.ibb.co (host gratuito de
terceiros), o que e um ponto unico de falha - se o link cair, a foto some em
todos os graficos/cards ao mesmo tempo. Este script baixa cada foto, redimensiona
para um thumbnail leve e salva em dashboard/photos/<apelido>.jpg, para servir
localmente via GitHub Pages junto com o resto do site.

Nao roda dentro do ambiente do Claude (o proxy da sandbox bloqueia i.ibb.co).
Rode isso na sua maquina, com internet normal:

    pip install requests pillow
    python3 fetch_photos.py <caminho-para-o-xlsx>

Depois rode build_data.py de novo - ele vai preferir o arquivo local se existir
em dashboard/photos/<apelido>.jpg.
"""
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

try:
    import requests
    from PIL import Image
except ImportError:
    print("Faltam dependencias. Rode: pip install requests pillow")
    sys.exit(1)

OUT_DIR = Path(__file__).parent / 'photos'
MAX_SIZE = 240  # px, lado maior


def slug(name):
    s = unicodedata.normalize('NFKD', name).encode('ascii', 'ignore').decode()
    s = re.sub(r'[^a-zA-Z0-9]+', '-', s).strip('-').lower()
    return s


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    OUT_DIR.mkdir(exist_ok=True)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    ws = wb['Banco de atletas']
    ok, skipped, failed = 0, 0, 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if not row[0]:
            continue
        nickname, photo_url = row[1], row[9]
        if not photo_url or not str(photo_url).startswith('http') or 'drive.google.com' in str(photo_url):
            print(f"  [skip] {nickname}: sem link direto de imagem ({photo_url})")
            skipped += 1
            continue
        dest = OUT_DIR / f"{slug(nickname)}.jpg"
        try:
            r = requests.get(photo_url, timeout=20)
            r.raise_for_status()
            tmp = dest.with_suffix('.tmp')
            tmp.write_bytes(r.content)
            img = Image.open(tmp).convert('RGB')
            img.thumbnail((MAX_SIZE, MAX_SIZE))
            img.save(dest, 'JPEG', quality=85)
            tmp.unlink()
            print(f"  [ok] {nickname} -> {dest.relative_to(Path(__file__).parent)}")
            ok += 1
        except Exception as e:
            print(f"  [FALHOU] {nickname}: {e}")
            failed += 1
    print(f"\nOK={ok} skip={skipped} falhou={failed}")
    if skipped:
        print("Atletas sem link direto (ex.: link do Google Drive) precisam de foto manual em dashboard/photos/.")


if __name__ == '__main__':
    main()
